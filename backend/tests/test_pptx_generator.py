import io

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.services import image_fetcher
from app.services import pptx_generator
from app.services.pptx_generator import build_pptx


def _jpeg_bytes():
    buf = io.BytesIO()
    Image.new("RGB", (160, 90), "#336699").save(buf, format="JPEG")
    return buf.getvalue()


def _deck():
    return [
        {"slide_number": 1, "title": "Cover", "bullets": ["Cover line."], "type": "cover"},
        {"slide_number": 2, "title": "Outline", "bullets": ["1. Item"], "type": "outline"},
        {
            "slide_number": 3,
            "title": "Identity Governance",
            "subtitle": "Key Theoretical Concepts",
            "bullets": ["Generated line one.", "Generated line two.", "Generated line three."],
            "type": "theory",
            "image_url": "https://example.invalid/photo.jpg",
            "image_keywords": ["identity", "governance"],
            "image_content": ["Generated line one.", "Generated line two."],
        },
        {"slide_number": 4, "title": "Thank You", "bullets": ["Questions"], "type": "end"},
    ]


def _render_text(monkeypatch, tmp_path):
    monkeypatch.setattr(pptx_generator, "fetch_image", lambda *a, **k: _jpeg_bytes())
    monkeypatch.setattr(image_fetcher, "_fetch_image_uncached", lambda *a, **k: _jpeg_bytes())
    image_fetcher._IMAGE_CACHE.clear()
    out = build_pptx(_deck(), 1, tmp_path / "deck.pptx", palette_id=1)
    image_fetcher._IMAGE_CACHE.clear()
    prs = Presentation(out)
    return [
        " ".join(
            shape.text_frame.text
            for shape in slide.shapes
            if shape.has_text_frame
        )
        for slide in prs.slides
    ]


def test_image_slide_still_renders_its_text(monkeypatch, tmp_path):
    """An image panel takes the right third; the generated text must survive."""
    text = _render_text(monkeypatch, tmp_path)
    assert "Generated line one." in text[2]
    assert "Generated line two." in text[2]
    assert "Generated line three." in text[2]


def test_image_prefetch_runs_before_rendering(monkeypatch, tmp_path):
    calls = []

    def fake_uncached(*args, **kwargs):
        calls.append(args[0])
        return _jpeg_bytes()

    monkeypatch.setattr(image_fetcher, "_fetch_image_uncached", fake_uncached)
    image_fetcher._IMAGE_CACHE.clear()
    build_pptx(_deck(), 1, tmp_path / "deck.pptx", palette_id=1)
    image_fetcher._IMAGE_CACHE.clear()
    assert calls == ["identity governance"]


def test_promote_text_dense_splits_content_and_image():
    slide = {
        "title": "Decision Flow",
        "type": "theory",
        "bullets": ["Step one.", "Step two."],
        "image_url": "document:diagram",
        "image_local_path": "/tmp/fig.png",
        "image_promote": True,
        "image_caption": "Figure 1: the flow",
    }
    out = pptx_generator._promote_text_dense_images([slide], {})
    assert len(out) == 2
    assert out[0]["type"] == "theory"
    assert "image_url" not in out[0]
    assert out[0]["bullets"] == ["Step one.", "Step two."]
    assert out[1]["type"] == "image"
    assert out[1]["image_url"] == "document:diagram"
    assert out[1]["bullets"] == []
    assert out[1]["subtitle"] == "Figure 1: the flow"


def test_generated_infographic_grows_the_deck(monkeypatch, tmp_path):
    monkeypatch.setattr(pptx_generator, "fetch_image", lambda *a, **k: _jpeg_bytes())
    monkeypatch.setattr(image_fetcher, "_fetch_image_uncached", lambda *a, **k: _jpeg_bytes())
    monkeypatch.setattr(pptx_generator, "image_is_generated", lambda *a, **k: True)
    image_fetcher._IMAGE_CACHE.clear()
    out = build_pptx(_deck(), 1, tmp_path / "deck.pptx", palette_id=1)
    image_fetcher._IMAGE_CACHE.clear()
    prs = Presentation(out)

    def slide_text(index):
        return " ".join(
            shape.text_frame.text
            for shape in prs.slides[index].shapes
            if shape.has_text_frame
        )

    assert len(prs.slides) == 5
    # The content slide keeps its bullets; the dedicated image slide does not.
    assert "Generated line three." in slide_text(2)
    assert "Generated line one." not in slide_text(3)
    assert len([sh for sh in prs.slides[3].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]) == 1


def test_planner_diagram_is_promoted_and_fills_its_own_slide(monkeypatch, tmp_path):
    def boom(*args, **kwargs):
        raise AssertionError("a planner diagram must not be classified via a photo lookup")

    monkeypatch.setattr(pptx_generator, "image_is_generated", boom)
    monkeypatch.setattr(pptx_generator, "fetch_image", lambda *a, **k: _jpeg_bytes())
    monkeypatch.setattr(image_fetcher, "_fetch_image_uncached", lambda *a, **k: _jpeg_bytes())
    image_fetcher._IMAGE_CACHE.clear()
    deck = [
        {"slide_number": 1, "title": "Cover", "bullets": ["Cover line."], "type": "cover"},
        {"slide_number": 2, "title": "Outline", "bullets": ["1. Item"], "type": "outline"},
        {
            "slide_number": 3,
            "title": "Approval Flow",
            "bullets": ["Draft the request.", "Route for review."],
            "type": "theory",
            "image_url": "https://example.invalid/diagram.jpg",
            "image_keywords": ["approval", "flow"],
            "image_content": ["Draft the request.", "Route for review."],
            "image_diagram": {"title": "Approval Flow", "steps": ["Draft", "Review", "Approve"]},
            "image_promote": True,
        },
        {"slide_number": 4, "title": "Thank You", "bullets": ["Questions"], "type": "end"},
    ]
    out = build_pptx(deck, 1, tmp_path / "deck.pptx", palette_id=1)
    image_fetcher._IMAGE_CACHE.clear()
    prs = Presentation(out)
    assert len(prs.slides) == 5
    text = " ".join(
        shape.text_frame.text for shape in prs.slides[2].shapes if shape.has_text_frame
    )
    assert "Draft the request." in text
    pics = [sh for sh in prs.slides[3].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pics) == 1
    assert pics[0].width == prs.slide_width
    assert pics[0].height == prs.slide_height


def test_promoted_diagram_fills_the_slide_without_cropping(monkeypatch, tmp_path):
    def slide_bytes():
        buf = io.BytesIO()
        Image.new("RGB", (1600, 900), "#336699").save(buf, format="JPEG")
        return buf.getvalue()

    monkeypatch.setattr(pptx_generator, "fetch_image", lambda *a, **k: slide_bytes())
    monkeypatch.setattr(image_fetcher, "_fetch_image_uncached", lambda *a, **k: slide_bytes())
    monkeypatch.setattr(pptx_generator, "image_is_generated", lambda *a, **k: True)
    image_fetcher._IMAGE_CACHE.clear()
    out = build_pptx(_deck(), 1, tmp_path / "deck.pptx", palette_id=1)
    image_fetcher._IMAGE_CACHE.clear()
    prs = Presentation(out)
    pics = [sh for sh in prs.slides[3].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pics) == 1
    pic = pics[0]
    # Generated 16:9 infographic on a 16:9 slide: edge to edge, uncropped.
    assert pic.width == prs.slide_width
    assert pic.height == prs.slide_height


def test_continuation_outlines_share_the_outline_block():
    slides = [
        {"type": "cover"},
        {"type": "outline"},
        {"type": "outline"},
        {"type": "outline"},
        {"type": "theory"},
    ]
    cover = {"type": "cover"}
    outline = {"type": "outline", "backgroundColor": "#111111"}
    content = {"type": "content", "backgroundColor": "#222222"}
    blocks = pptx_generator._match_blocks(slides, [cover, outline, content])
    assert blocks[1] is outline
    assert blocks[2] is outline
    assert blocks[3] is outline
    # The extra outline pages must not consume the content block.
    assert blocks[4] is content


def test_outline_numbering_continues_across_rendered_pages():
    outline_pages = [
        {
            "slide_number": 2,
            "title": "Presentation Outline",
            "type": "outline",
            "bullets": ["**Opening**", "1. Deck", "**Core Topics**", "2. First", "3. Second"],
        },
        {
            "slide_number": 3,
            "title": "Presentation Outline (continued)",
            "type": "outline",
            "bullets": ["4. Third", "**Closing**", "5. Key Takeaways & Next Steps"],
        },
    ]
    prs = Presentation()
    prs.slide_width = pptx_generator.Inches(pptx_generator.SLIDE_W)
    prs.slide_height = pptx_generator.Inches(pptx_generator.SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pptx_generator._render_outline(prs, slide, pptx_generator.PALETTES[1], {}, outline_pages[1])
    text = slide.shapes[-1].text_frame.text
    assert "04" in text
    assert "05" in text
    assert "01" not in text
